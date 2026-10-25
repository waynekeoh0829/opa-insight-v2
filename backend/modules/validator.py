import pandas as pd, os, re

REQUIRED = {"order_id","product_name","quantity","revenue","gst","date"}
ALIASES  = {
    "order id":"order_id","orderid":"order_id","order no":"order_id","order number":"order_id",
    "product":"product_name","item":"product_name","product name":"product_name","item name":"product_name","name":"product_name",
    "qty":"quantity","amount":"quantity","units":"quantity",
    "total":"revenue","sales":"revenue","price":"revenue","revenue":"revenue","total revenue":"revenue",
    "tax":"gst","vat":"gst","gst amount":"gst","gst":"gst",
    "date":"date","order date":"date","transaction date":"date","purchase date":"date","created":"date","created at":"date"
}

def _normalise_cols(df):
    renamed = {}
    for col in df.columns:
        plain = col.strip().lower()
        norm = plain.replace("-","_").replace(" ","_")
        if norm in REQUIRED: renamed[col]=norm
        elif plain in ALIASES: renamed[col]=ALIASES[plain]
    return df.rename(columns=renamed)

def _read_docx(path):
    from docx import Document
    doc = Document(path)
    rows = []
    for table in doc.tables:
        if not table.rows: continue
        headers = [c.text.strip() for c in table.rows[0].cells]
        if len(headers) < 4: continue
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            rows.append(dict(zip(headers, cells)))
    if rows:
        return pd.DataFrame(rows)
    lines=[p.text.strip() for p in doc.paragraphs if p.text.strip()]
    data=[]
    for line in lines:
        parts=re.split(r'[,|\t]+',line)
        if len(parts)>=4: data.append(parts)
    if not data: return pd.DataFrame()
    return pd.DataFrame(data[1:],columns=data[0])

def validate_file(path, ext):
    try:
        if ext=="csv":
            df=None
            for enc in ["utf-8","utf-8-sig","latin-1","cp1252"]:
                try: df=pd.read_csv(path,encoding=enc); break
                except: pass
        elif ext in ("xlsx","xls"): df=pd.read_excel(path)
        elif ext in ("docx","doc"): df=_read_docx(path)
        else: return {"ok":False,"error":f"Unsupported: .{ext}"}
        if df is None or df.empty: return {"ok":False,"error":"File empty or unreadable"}
        df=_normalise_cols(df)
        missing=REQUIRED-set(df.columns)
        if missing: return {"ok":False,"error":f"Missing columns: {', '.join(sorted(missing))}","found":list(df.columns)}
        preview=df.head(5).fillna("").astype(str).to_dict(orient="records")
        return {"ok":True,"columns":list(df.columns),"row_count":len(df),"preview":preview,"df":df}
    except Exception as e:
        return {"ok":False,"error":str(e)}
